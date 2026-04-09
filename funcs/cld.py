# Copyright (c) 2025 Jeroen F. Uleman. Licensed under CC BY-NC 4.0.
# Non-commercial use only. See LICENSE for details.
import os
import pandas as pd
import numpy as np
import warnings
import networkx as nx
from types import SimpleNamespace
from openpyxl import load_workbook

class Extract:
    def __init__(self, file_path):
        self.file_path = file_path
        self.variables = []
        self.var_to_type = {}
        self.adjacency_matrix = None
        self.interactions_matrix = None
        self.lb_matrix = None   # per-connection lower bounds (NaN = use default sampling)
        self.ub_matrix = None   # per-connection upper bounds (NaN = use default sampling)

    def extract_settings(self, double_factor_interventions_setting=None):
        """Extract all settings from the Kumu Excel file."""
        self.adjacency_matrix_from_kumu()

        if np.abs(self.interactions_matrix).sum() > 0:  # Interaction terms specified
            print("Solving an SDM with interaction terms.")
            print("By default, only single interventions will be simulated. "
                  "To simulate interventions on two variables simultaneously, "
                  "set s.double_factor_interventions = True")
            s = SimpleNamespace(**{"interaction_terms": 1,
                                   "solve_analytically": 0,
                                   "double_factor_interventions": 0})
        else:
            print("No interaction terms specified so will solve linear SDM.")
            s = SimpleNamespace(**{"interaction_terms": 0,
                                   "solve_analytically": 1,
                                   "double_factor_interventions": 0})

        if double_factor_interventions_setting is not None:
            s.double_factor_interventions = double_factor_interventions_setting

        if s.double_factor_interventions and s.interaction_terms == False:
            warnings.warn("Without interaction terms, double factor interventions are not meaningful. "
                          "Consider setting double_factor_interventions to False.")

        s.stocks = [var for var in self.variables if self.var_to_type[var].lower() == 'stock']
        s.auxiliaries = [var for var in self.variables if self.var_to_type[var].lower() == 'auxiliary']
        s.constants = [var for var in self.variables if self.var_to_type[var].lower() == 'constant']
        s.variables = list(self.variables)
        s.stocks_and_constants = [var for var in self.variables if self.var_to_type[var] in ['stock', 'constant']]
        s.stocks_and_auxiliaries = [var for var in self.variables if self.var_to_type[var] in ['stock', 'auxiliary']]
        s.var_to_type = {var: self.var_to_type[var] for var in self.variables}

        s.variable_of_interest = list(self.variable_of_interest)
        s.centrality = self.centrality
        s.intervention_variables = self.intervention_variables
        s.intervention_strengths = self.intervention_strengths

        if len(s.intervention_variables) == 0:
            raise Exception("There should be at least one intervention variable specified in the Excel file.")

        if s.double_factor_interventions:
            double_intervention_variables = []
            for i, var in enumerate(s.intervention_variables):
                for j in range(i + 1, len(s.intervention_variables)):
                    double_intervention_variables.append(var + '+' + s.intervention_variables[j])
            s.intervention_variables += double_intervention_variables

        self.df_adj.rename(index=dict(zip(self.variables, s.variables)),
                           columns=dict(zip(self.variables, s.variables)), inplace=True)

        s.df_adj = self.df_adj
        s.interactions_matrix = self.interactions_matrix
        s.lb_matrix = self.lb_matrix
        s.ub_matrix = self.ub_matrix

        # Extend adjacency matrix to include interaction-term edges (for loop detection)
        s.df_adj_incl_interactions = s.df_adj.copy()
        to_list, from1_list, from2_list = np.nonzero(s.interactions_matrix)
        for i in range(int(np.abs(s.interactions_matrix).sum())):
            to, from1, from2 = to_list[i], from1_list[i], from2_list[i]
            value = s.interactions_matrix[to, from1, from2]
            s.df_adj_incl_interactions.iloc[to, from1] = value
            s.df_adj_incl_interactions.iloc[to, from2] = value

        self.s = s
        return s

    def check_loops(self, df_e, df_c):
        """Check whether all loops contain a stock and report the balancing/reinforcing ratio."""
        stocks = [var for var in self.variables if self.var_to_type[var].lower() == "stock"]
        num_stocks_and_auxiliaries = len([var for var in self.variables
                                         if self.var_to_type[var].lower() in ["stock", "auxiliary"]])

        max_loops_check = min(5, num_stocks_and_auxiliaries)

        if (self.df_adj == -999).any().any() > 0:
            temp_df = self.df_adj.copy()
            temp_df = temp_df.replace(-999, 1)

        G = nx.from_numpy_array(np.array(self.df_adj).T, create_using=nx.DiGraph)
        var_names = list(self.df_adj.columns)
        G = nx.relabel_nodes(G, dict(enumerate(var_names)))
        feedback_loops = list(nx.simple_cycles(G, length_bound=max_loops_check))
        num_loops = len(feedback_loops)

        if num_loops > 0:
            print(f"\n{num_loops} feedback loops of maximum length {max_loops_check}")

            loops_wo_stocks = [loop for loop in feedback_loops
                               if sum(1 for x in loop if x in stocks) == 0]

            if loops_wo_stocks:
                print(len(loops_wo_stocks), "loops do not have a stock, which is",
                      round(len(loops_wo_stocks) / num_loops * 100, 5), "% of all loops")
                print("Loops without stocks:", loops_wo_stocks)
                raise Exception("All loops should have at least one stock, redo the labeling")
            else:
                print("All loops have at least one stock")

            if (self.df_adj == -999).any().any() > 0:
                num_balancing = 0
                for loop in feedback_loops:
                    num_min = 0
                    loop_closed = loop + [loop[0]]
                    for i in range(len(loop_closed) - 1):
                        pol = df_c.loc[
                            ((df_c.From == loop_closed[i]) * 1 + (df_c.To == loop_closed[i + 1]) * 1) == 2,
                            "Type"
                        ].values[0]
                        if str(pol) == "-":
                            num_min += 1
                    if num_min % 2 != 0:
                        num_balancing += 1

                print(f"{num_balancing} ({round(num_balancing / num_loops * 100, 2)}%) "
                      f"of these loops are balancing loops")
                if max_loops_check == num_stocks_and_auxiliaries:
                    print("The max length of loops checked equals the number of stocks and auxiliaries; "
                          "all loops are considered\n")
                else:
                    print("The max length of loops checked is smaller than the number of stocks and auxiliaries; "
                          "there may be more loops in the CLD\n")
        else:
            print("No feedback loops found in the CLD")

        bc = nx.betweenness_centrality(G, k=None, normalized=True)
        cc = nx.closeness_centrality(G.reverse(), wf_improved=False)
        self.centrality = {
            "betweenness": dict(sorted(bc.items(), key=lambda x: x[1], reverse=True)),
            "closeness":   dict(sorted(cc.items(), key=lambda x: x[1], reverse=True)),
        }

    def extract_adjacency_matrix(self):
        """Extract the adjacency matrix from an Excel file exported from Kumu (Kumu.io).

        The file must contain:
          - 'Elements' sheet: variables with columns Label, Type, Tags, Description
          - 'Connections' sheet: causal links with columns From, Type, To
                                 (optional LB and UB columns for per-connection sampling bounds)
          - 'Interaction terms' sheet (optional): second-order terms with From1, From2, Type, To
        """
        df_e = pd.read_excel(self.file_path, sheet_name="Elements")
        df_c = pd.read_excel(self.file_path, sheet_name="Connections")

        df_e = df_e[["Label", "Type", "Tags", "Description"]]
        conn_cols = ["From", "Type", "To"]
        if "LB" in df_c.columns:
            conn_cols.append("LB")
        if "UB" in df_c.columns:
            conn_cols.append("UB")
        df_c = df_c[conn_cols]
        has_custom_bounds_col = "LB" in df_c.columns and "UB" in df_c.columns

        self.original_variables = list(df_e["Label"])

        for var in self.original_variables:
            if "+" in var:
                raise Exception(f'Variable name {var} contains a disallowed special character (+).')
            if "*" in var:
                raise Exception(f'Variable name {var} contains a disallowed special character (*).')

        self.variables = [" ".join(var.split()) for var in self.original_variables]
        self.var_to_type = dict(zip(self.variables, list(df_e["Type"])))
        self.original_to_cleaned_var = dict(zip(self.original_variables, self.variables))

        for var in self.original_variables:
            cleaned = self.original_to_cleaned_var[var]
            type_ = str(self.var_to_type[cleaned]).lower()
            if type_ not in ['stock', 'auxiliary', 'constant']:
                if var in list(df_c['To']):
                    print(f'Warning: {cleaned} has no known label (got "{type_}"), '
                          f'defaulting to stock (has incoming links).')
                    self.var_to_type[cleaned] = 'stock'
                else:
                    print(f'Warning: {cleaned} has no known label (got "{type_}"), '
                          f'defaulting to constant (no incoming links).')
                    self.var_to_type[cleaned] = 'constant'

        self.intervention_variables = [
            self.original_to_cleaned_var[var]
            for var in list(df_e.loc[df_e["Tags"] != 0, "Label"])
        ]
        self.intervention_strengths = dict(zip(self.variables, list(df_e["Tags"])))
        self.variable_of_interest = list(df_e.loc[df_e["Description"] == "VOI", "Label"])

        if len(self.variable_of_interest) == 1:
            print("Variable of interest:", self.variable_of_interest[0])
        else:
            print("Variables of interest:", self.variable_of_interest)
        print("with", len(self.intervention_variables), "intervention variables")

        num_variables = len(self.variables)
        self.adjacency_matrix = np.zeros((num_variables, num_variables))
        self.lb_matrix = np.full((num_variables, num_variables), np.nan)
        self.ub_matrix = np.full((num_variables, num_variables), np.nan)

        for i, origin in enumerate(df_c["From"]):
            if origin not in self.original_variables:
                raise Exception(f'Origin variable "{origin}" in Connections not found in Elements.')
            destination = df_c["To"][i]
            if destination not in self.original_variables:
                raise Exception(f'Destination variable "{destination}" in Connections not found in Elements.')

            temp = df_c["Type"][i]
            if str(temp) == '+':
                polarity = 1
            elif str(temp) == '-':
                polarity = -1
            else:
                polarity = -999

            oi = self.original_variables.index(origin)
            di = self.original_variables.index(destination)
            self.adjacency_matrix[di, oi] = polarity

            if has_custom_bounds_col:
                lb_val = df_c["LB"][i]
                ub_val = df_c["UB"][i]
                if pd.notna(lb_val) and pd.notna(ub_val):
                    if float(lb_val) >= float(ub_val):
                        raise ValueError(
                            f"LB ({lb_val}) must be strictly less than UB ({ub_val}) "
                            f"for connection {origin} → {destination}."
                        )
                    self.lb_matrix[di, oi] = float(lb_val)
                    self.ub_matrix[di, oi] = float(ub_val)

        self.df_adj = pd.DataFrame(self.adjacency_matrix,
                                   columns=self.variables,
                                   index=self.variables)

        constants = [var for var in self.variables if self.var_to_type[var].lower() == 'constant']

        for const in constants:
            num_incoming = np.sum(np.abs(self.df_adj.loc[const, :]))
            if num_incoming != 0:
                print(f'Removed {num_incoming} incoming link(s) for constant {const}')
                self.df_adj.loc[const, :] = 0

        for var in self.variables:
            if var not in constants:
                if np.sum(np.abs(self.df_adj.loc[var, :])) == 0:
                    raise Exception(f'Non-constant variable "{var}" has no incoming links.')

        self.check_loops(df_e, df_c)

    def extract_interactions_matrix(self):
        """Extract the interactions matrix from the 'Interaction terms' sheet in the Kumu Excel file."""
        wb = load_workbook(self.file_path, read_only=True)

        num_variables = len(self.variables)
        self.interactions_matrix = np.zeros((num_variables, num_variables, num_variables))

        if 'Interaction terms' in wb.sheetnames:
            df_i = pd.read_excel(self.file_path, sheet_name="Interaction terms")
            df_i = df_i[["From1", "From2", "Type", "To"]]

            for i, origin_1 in enumerate(df_i["From1"]):
                origin_2 = df_i["From2"][i]
                destination = df_i["To"][i]

                temp = df_i["Type"][i]
                if str(temp) == '+':
                    polarity = 1
                elif str(temp) == '-':
                    polarity = -1
                else:
                    polarity = -999

                origin_1_index = self.original_variables.index(origin_1)
                origin_2_index = self.original_variables.index(origin_2)
                destination_index = self.original_variables.index(destination)

                self.interactions_matrix[destination_index, origin_2_index, origin_1_index] = polarity

    def adjacency_matrix_from_kumu(self):
        """Extract the adjacency matrix and interactions matrix from the Kumu Excel file."""
        self.extract_adjacency_matrix()
        self.extract_interactions_matrix()


### TESTING ###
    def test_extraction(self):
        """Test the CLD extraction by creating an exemplar Kumu table and comparing the results."""
        data = {
            "From": ["A", "B", "C"],
            "Type": ["+", "-", "+"],
            "To": ["B", "C", "A"]
        }
        data_int = {
            "From1": ["A", "B"],
            "From2": ["C", "C"],
            "Type": ["+", "+"],
            "To": ["B", "A"]
        }

        df_e = pd.DataFrame(data["From"], columns=["Label"])
        df_e["Type"] = ["stock", "auxiliary", "constant"]
        df_e["Tags"] = [0, -1, 1]
        df_e["Description"] = ["VOI", None, None]
        df_c = pd.DataFrame(data)
        df_i = pd.DataFrame(data_int)

        original_file_path = self.file_path
        test_file_path = os.path.join(os.path.dirname(__file__), '..', 'tests', 'evidence_table.xlsx')
        os.makedirs(os.path.dirname(test_file_path), exist_ok=True)

        with pd.ExcelWriter(test_file_path) as writer:
            df_e.to_excel(writer, sheet_name='Elements', index=False)
            df_c.to_excel(writer, sheet_name='Connections', index=False)
            df_i.to_excel(writer, sheet_name='Interaction terms', index=False)

        self.file_path = test_file_path
        self.adjacency_matrix_from_kumu()
        self.file_path = original_file_path

        # Note: C is a constant so its incoming link (B→C) is automatically removed.
        expected_adjacency_matrix = np.array([[0, 0, 1],
                                              [1, 0, 0],
                                              [0, 0, 0]])

        expected_interactions_matrix = np.array([[[0, 0, 0],
                                                  [0, 0, 0],
                                                  [0, 1, 0]],
                                                 [[0, 0, 0],
                                                  [0, 0, 0],
                                                  [1, 0, 0]],
                                                 [[0, 0, 0],
                                                  [0, 0, 0],
                                                  [0, 0, 0]]])

        assert np.all(expected_adjacency_matrix == self.adjacency_matrix)
        assert np.all(expected_interactions_matrix == self.interactions_matrix)
        assert np.all([x in self.variables for x in data["From"]])
        assert np.all([x in data["From"] for x in self.variables])
        print("Test for loading KUMU table passed.")
